"""文件预览（2026-09-30）：对话里的附件点开就在 app 里看，不再跳浏览器。

GET /api/files/{fid}/preview 回这个文件怎么显示（view）和显示用的内容：
- image：图片。原图用 /api/files/{fid}?preview=1 取（网页认不得的格式、特别大的图转成 JPEG，见 files.preview_image）。
- pages：PDF，和 PyMuPDF 打得开的 EPUB / XPS / CBZ / FB2 / MOBI。回每页宽高；页图 GET /api/files/{fid}/page/{n}?w=1200
  （宽度就近取 800 / 1200 / 1600 / 2000，渲染一次存在原件旁边，和缩略图一样）。
- doc：一串块，{"md": Markdown} 或 {"table": 表格}。Word（标题、列表、粗体、表格）、PowerPoint（每页一节，带表格和备注）、
  Excel / CSV / TSV（每个工作表一张表）、Markdown、Jupyter 笔记本。长文按段切成几块，app 一块一块排，不会一次卡住。
- text：纯文字。代码、JSON（排好缩进）、日志、压缩包的文件列表用等宽（mono）；.txt、网页抽出的正文不用。
- audio / video：播放；音频带上传时的转写。
- none：看不了，只有文件信息（app 给「用浏览器打开」）。
只读原件，不改库。PyMuPDF 没装时 PDF 退回上传时抽的文字。PyMuPDF 不能多线程同时用，所有调用都过 _MU 锁。
"""
from __future__ import annotations

import csv
import io
import json
import re
import threading
import uuid
import zipfile
from datetime import date, datetime, time
from html.parser import HTMLParser
from pathlib import Path

from fastapi import APIRouter, HTTPException
from fastapi.responses import FileResponse

from chat import _lock
from files import IMMUTABLE as CACHE, adb, count, summary
from i18n import L

router = APIRouter()

TEXT_MAX = 200_000          # 纯文字最多给这么多字
MD_MAX = 80_000             # 要排版的（Markdown、Word、PPT）少给一些
MD_CHUNK = 4_000            # 按段切块，每块大约这么长
TABLE_ROWS = 200            # 每张表最多这么多行、列，每格这么多字
TABLE_COLS = 26
CELL_MAX = 200
CELLS_MAX = 12_000          # 一个文件里所有表加起来最多这么多格
PAGES_MAX = 500             # 页数再多，宽高只列前 500 页
PAGE_WIDTHS = (800, 1200, 1600, 2000)
PAGE_PIXELS = 8_000_000     # 一页图最多 800 万像素：特别长的页按面积缩
PAGED_EXT = {".pdf", ".epub", ".xps", ".oxps", ".cbz", ".fb2", ".mobi", ".svg"}   # SVG 也当一页渲染成图（不在 app 里跑它的脚本）
PROSE_EXT = {".txt", ".text", ".rtf", ".html", ".htm"}
_MU = threading.Lock()


def _row(fid: str):
    with _lock, adb() as conn:
        r = conn.execute("SELECT * FROM attachments WHERE id=?", (fid,)).fetchone()
    if not r or not Path(r["path"]).is_file():
        raise HTTPException(404, L("文件不在了", "File no longer exists"))
    return r


# —— 小工具 ——

def _decode(raw: bytes) -> str:
    for enc in ("utf-8-sig", "gb18030"):   # Excel 导出的中文 CSV 常是 GBK
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", "replace")


def _binary(raw: bytes) -> bool:
    return b"\x00" in raw[:8192]


_MD_CHARS = re.compile(r"([\\`*_\[\]])")
_MD_LEAD = re.compile(r"^(\s*)([#>+-]|\d+[.)])(?=\s)", re.M)


def _esc(s: str) -> str:
    """Word / PPT 里的字原样显示：不让里面的 * _ # 1. 被当成 Markdown。"""
    return _MD_LEAD.sub(r"\1\\\2", _MD_CHARS.sub(r"\\\1", s))


def _chunks(md: str) -> list[dict]:
    """长 Markdown 在空行处切成几块；代码块（```）中间不切。"""
    out, buf, size, fence = [], [], 0, False
    for para in md.split("\n\n"):
        buf.append(para)
        size += len(para)
        fence ^= para.count("```") % 2 == 1
        if size >= MD_CHUNK and not fence:
            out.append({"md": "\n\n".join(buf)})
            buf, size = [], 0
    if buf and "".join(buf).strip():
        out.append({"md": "\n\n".join(buf)})
    return out


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float):
        return f"{v:.10g}"
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M") if (v.hour, v.minute, v.second) != (0, 0, 0) else v.strftime("%Y-%m-%d")
    if isinstance(v, (date, time)):
        return v.isoformat()
    s = str(v).strip()
    return s if len(s) <= CELL_MAX else s[: CELL_MAX - 1] + "…"


class _Budget:
    """一个文件里的表格总格数。"""
    def __init__(self) -> None:
        self.left = CELLS_MAX


def _table(rows: list[list[str]], budget: _Budget, rows_total: int | None = None, cols_total: int | None = None) -> dict | None:
    rows = [[_cell(c) for c in r[:TABLE_COLS]] for r in rows[:TABLE_ROWS]]
    while rows and not any(rows[-1]):          # 去掉尾巴上的空行、空列
        rows.pop()
    width = max((max((i + 1 for i, c in enumerate(r) if c), default=0) for r in rows), default=0)
    if not width:
        return None
    rows = [(r + [""] * width)[:width] for r in rows]
    fit = max(1, budget.left // width)
    rows_total = max(rows_total or 0, len(rows))
    rows = rows[:fit]
    budget.left -= len(rows) * width
    return {"rows": rows, "head": True, "rowsTotal": rows_total, "colsTotal": max(cols_total or 0, width)}


# —— 各种文件 ——

def _mupdf():
    try:
        import pymupdf
        return pymupdf
    except ImportError:
        try:
            import fitz  # 老版本的包名
            return fitz
        except ImportError:
            return None


def _open_paged(m, path: Path):
    doc = m.open(str(path))
    if doc.is_reflowable:   # EPUB 这类没有固定页面：按手机屏排版（两个接口排法一样，页码才对得上）
        doc.layout(width=400, height=620, fontsize=12)
    return doc


def _pages(path: Path) -> dict | None:
    m = _mupdf()
    if not m:
        return None
    with _MU:
        try:
            doc = _open_paged(m, path)
        except Exception:  # noqa: BLE001 — 打不开就退回文字
            return None
        try:
            if doc.needs_pass:
                return {"view": "none", "note": L("这个文件有密码，预览不了。", "This file is password-protected, so it can't be previewed.")}
            n = doc.page_count
            pages = []
            for i in range(min(n, PAGES_MAX)):
                r = doc[i].rect
                pages.append([round(r.width, 1), round(r.height, 1)])
        finally:
            doc.close()
    if not pages:
        return None
    return {"view": "pages", "pageCount": n, "pages": pages}


def _para(p) -> str:
    """Word 的一段 → 一行 Markdown：标题、列表、引用、粗体 / 斜体。"""
    style = (p.style.name if p.style is not None else "") or ""
    try:
        parts = list(p.iter_inner_content())   # python-docx 1.x：连超链接里的字一起
    except AttributeError:
        parts = list(p.runs)
    spans: list[list] = []   # [粗, 斜, 字]，相邻同样式的合在一起
    for x in parts:
        text = x.text or ""
        if not text:
            continue
        bold, ital = bool(getattr(x, "bold", False)), bool(getattr(x, "italic", False))
        if spans and spans[-1][0] == bold and spans[-1][1] == ital:
            spans[-1][2] += text
        else:
            spans.append([bold, ital, text])
    plain = "".join(s[2] for s in spans).strip()
    if not plain:
        return ""
    m = re.match(r"(?:Heading|标题)\s*(\d)", style)
    if m or style in ("Title", "标题"):
        level = min(int(m.group(1)), 4) if m else 1
        return "#" * level + " " + _esc(plain).replace("\n", " ")
    body = ""
    for bold, ital, text in spans:
        lead, core, tail = re.match(r"^(\s*)(.*?)(\s*)$", text, re.S).groups()
        core = _esc(core)
        if core and bold:
            core = f"**{core}**"
        if core and ital:
            core = f"*{core}*"
        body += lead + core + tail
    body = body.strip()
    ppr = p._p.pPr
    num = ppr.numPr if ppr is not None else None
    if num is not None or re.match(r"List (Bullet|Number)", style):   # 编号可能写在段落上，也可能写在样式里
        level = num.ilvl.val if num is not None and num.ilvl is not None else 0
        return "  " * min(level, 4) + "- " + body.replace("\n", " ")
    if "Quote" in style:
        return "> " + body.replace("\n", "\n> ")
    return body


def _docx(path: Path) -> dict:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    d = docx.Document(str(path))
    budget, blocks, buf, size, cut = _Budget(), [], [], 0, False

    def flush() -> None:
        if buf:
            blocks.extend(_chunks("\n\n".join(buf)))
            buf.clear()

    for el in d.element.body.iterchildren():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            line = _para(Paragraph(el, d))
            if line:
                buf.append(line)
                size += len(line)
        elif tag == "tbl":
            rows = []
            for row in Table(el, d).rows:
                cells, seen = [], None
                for c in row.cells:   # 合并的格子 python-docx 会重复给，重复的留空、列对得上
                    cells.append("" if c._tc is seen else c.text.strip())
                    seen = c._tc
                rows.append(cells)
            tbl = _table(rows, budget)
            if tbl:
                flush()
                blocks.append({"table": tbl})
                size += sum(len(c) for r in tbl["rows"] for c in r)
        if size > MD_MAX:
            cut = True
            break
    flush()
    return {"view": "doc", "blocks": blocks, "truncated": cut, "label": "Word"}


def _shape_blocks(shapes, lines: list[str], blocks: list[dict], budget: _Budget, skip=None) -> None:
    from pptx.shapes.group import GroupShape
    for sh in sorted(shapes, key=lambda s: ((s.top or 0), (s.left or 0))):
        if skip is not None and sh.shape_id == skip:
            continue
        if isinstance(sh, GroupShape):
            _shape_blocks(sh.shapes, lines, blocks, budget)
        elif getattr(sh, "has_table", False) and sh.has_table:
            tbl = _table([[c.text.strip() for c in r.cells] for r in sh.table.rows], budget)
            if tbl:
                if lines:
                    blocks.append({"md": "\n\n".join(lines)})
                    lines.clear()
                blocks.append({"table": tbl})
        elif getattr(sh, "has_text_frame", False) and sh.has_text_frame:
            paras = [(p.level, "".join(r.text for r in p.runs).strip()) for p in sh.text_frame.paragraphs]
            paras = [(lv, x) for lv, x in paras if x]
            if len(paras) == 1 and paras[0][0] == 0:
                lines.append(_esc(paras[0][1]))
            else:
                lines.append("\n".join("  " * min(lv, 4) + "- " + _esc(x) for lv, x in paras))


def _pptx(path: Path) -> dict:
    from pptx import Presentation
    prs = Presentation(str(path))
    budget, blocks, size, cut, n = _Budget(), [], 0, False, 0
    for i, slide in enumerate(prs.slides, 1):
        n = i
        title_sh = slide.shapes.title
        title = title_sh.text_frame.text.strip() if title_sh is not None and title_sh.has_text_frame else ""
        head = "### " + L(f"第 {i} 页", f"Slide {i}") + (f" · {_esc(title)}" if title else "")
        lines: list[str] = [head]
        before = len(blocks)
        _shape_blocks(slide.shapes, lines, blocks, budget, skip=title_sh.shape_id if title_sh is not None else None)
        if slide.has_notes_slide:
            notes = (slide.notes_slide.notes_text_frame.text or "").strip() if slide.notes_slide.notes_text_frame is not None else ""
            if notes:
                lines.append("> " + L("备注：", "Notes: ") + _esc(notes).replace("\n", "\n> "))
        if lines:
            blocks.append({"md": "\n\n".join(lines)})
        size += sum(len(b["md"]) if "md" in b else 50 for b in blocks[before:])
        if size > MD_MAX:
            cut = True
            break
    total = len(prs.slides)
    return {"view": "doc", "blocks": blocks, "truncated": cut or n < total,
            "label": L(f"PowerPoint · {total} 页", f"PowerPoint · {count(total, 'slide')}")}


def _xlsx(path: Path) -> dict:
    import openpyxl
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    budget, blocks = _Budget(), []
    try:
        sheets = wb.worksheets
        for ws in sheets:
            try:
                if ws.calculate_dimension() == "A1:A1":   # 有的工具写错了尺寸，只读模式会只给第一格
                    ws.reset_dimensions()
            except ValueError:
                pass
            rows = [list(r) for _, r in zip(range(TABLE_ROWS), ws.iter_rows(values_only=True))]
            tbl = _table(rows, budget, ws.max_row, ws.max_column)
            if len(sheets) > 1:
                blocks.append({"md": "### " + _esc(ws.title)})
            blocks.append({"table": tbl} if tbl else {"md": L("（空）", "(Empty)")})
            if budget.left <= 0:
                break
    finally:
        wb.close()
    return {"view": "doc", "blocks": blocks, "label": L(f"Excel · {len(sheets)} 个工作表", f"Excel · {count(len(sheets), 'sheet')}")}


def _csv(path: Path, ext: str) -> dict:
    text = _decode(path.read_bytes()[:8_000_000])
    try:
        dialect = csv.excel_tab if ext == ".tsv" else csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows, total = [], 0
    for row in csv.reader(io.StringIO(text), dialect):
        total += 1
        if len(rows) < TABLE_ROWS:
            rows.append(row)
    tbl = _table(rows, _Budget(), total)
    return {"view": "doc", "blocks": [{"table": tbl}] if tbl else [], "label": ext[1:].upper()}


def _notebook(text: str) -> dict:
    nb = json.loads(text)
    lang = ((nb.get("metadata") or {}).get("kernelspec") or {}).get("language") or "python"
    src = lambda v: "".join(v) if isinstance(v, list) else str(v or "")  # noqa: E731
    parts = []
    for c in nb.get("cells") or []:
        body = src(c.get("source")).strip()
        if c.get("cell_type") == "markdown":
            parts.append(body)
        elif c.get("cell_type") == "code":
            parts.append(f"```{lang}\n{body}\n```")
            for o in c.get("outputs") or []:
                out = src(o.get("text") or (o.get("data") or {}).get("text/plain")).strip()
                if out:
                    parts.append("```\n" + out[:3000] + "\n```")
    md = "\n\n".join(p for p in parts if p)
    return {"view": "doc", "blocks": _chunks(md[:MD_MAX]), "truncated": len(md) > MD_MAX, "label": "Jupyter"}


class _HtmlText(HTMLParser):
    SKIP = {"script", "style", "noscript", "template", "head"}
    BLOCK = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "section", "article", "header", "footer",
             "blockquote", "pre", "table", "ul", "ol", "hr"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.out: list[str] = []
        self.skip = 0

    def handle_starttag(self, tag, attrs):  # noqa: ANN001
        if tag in self.SKIP:
            self.skip += 1
        elif tag in self.BLOCK:
            self.out.append("\n")

    def handle_endtag(self, tag):  # noqa: ANN001
        if tag in self.SKIP:
            self.skip = max(0, self.skip - 1)
        elif tag in self.BLOCK:
            self.out.append("\n")

    def handle_data(self, data):  # noqa: ANN001
        if not self.skip:
            self.out.append(data)


def _html_text(src: str) -> str:
    p = _HtmlText()
    p.feed(src)
    lines = [re.sub(r"[ \t\r\f\v ]+", " ", x).strip() for x in "".join(p.out).split("\n")]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def _text(text: str, mono: bool, label: str) -> dict:
    return {"view": "text", "text": text[:TEXT_MAX], "mono": mono, "truncated": len(text) > TEXT_MAX, "label": label}


def _zip(path: Path) -> dict:
    with zipfile.ZipFile(path) as z:
        infos = [i for i in z.infolist() if not i.is_dir()]
    lines = [f"{_zname(i)}  ·  {_size(i.file_size)}" for i in infos[:3000]]
    out = _text("\n".join(lines), True, L(f"压缩包 · {len(infos)} 个文件", f"Archive · {count(len(infos), 'file')}"))
    out["truncated"] = len(infos) > 3000
    return out


def _zname(i: zipfile.ZipInfo) -> str:
    """没标 UTF-8 的文件名 Python 按 CP437 解：Windows 打的包多半是 GBK，换回来。"""
    if i.flag_bits & 0x800:
        return i.filename
    raw = i.filename.encode("cp437", "replace")
    for enc in ("utf-8", "gb18030"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return i.filename


def _size(n: int) -> str:
    return f"{n / 1024 / 1024:.1f} MB" if n >= 1024 * 1024 else f"{n / 1024:.0f} KB" if n >= 1024 else f"{n} B"


def _image(path: Path) -> dict:
    out: dict = {"view": "image", "label": L("图片", "Image")}
    try:
        from PIL import Image
        try:
            import pillow_heif
            pillow_heif.register_heif_opener()
        except ImportError:
            pass
        with Image.open(path) as im:   # 只读文件头，不解码
            w, h = im.size
            if im.getexif().get(0x0112) in (5, 6, 7, 8):   # 竖着拍的：EXIF 说要转 90°
                w, h = h, w
        out.update(width=w, height=h, label=L(f"图片 · {w}×{h}", f"Image · {w}×{h}"))
    except Exception:  # noqa: BLE001 — 量不出尺寸也照样显示
        pass
    return out


def build(path: Path, name: str, mime: str, kind: str, stored_text: str | None) -> dict:
    ext = Path(name).suffix.lower() or path.suffix.lower()
    mime = (mime or "").lower()
    if kind == "image" and ext != ".svg" and mime != "image/svg+xml":
        return _image(path)
    if kind == "audio":
        return {"view": "audio", "transcript": stored_text or None, "label": L("音频", "Audio")}
    if kind == "video":
        return {"view": "video", "label": L("视频", "Video")}
    if ext in PAGED_EXT or mime in ("application/pdf", "image/svg+xml"):
        got = _pages(path)
        if got:
            if got["view"] == "pages":
                label = "PDF" if ext == ".pdf" or mime == "application/pdf" else ext[1:].upper()
                got["label"] = label if ext == ".svg" else L(f"{label} · {got['pageCount']} 页", f"{label} · {count(got['pageCount'], 'page')}")
            return got
        if stored_text:   # 服务器没装 PyMuPDF / 打不开：给上传时抽的文字
            out = _text(stored_text, False, "PDF")
            out["note"] = L("只能看抽出来的文字（服务器上装了 PyMuPDF 才能看原样的页面）。", "Text only (install PyMuPDF on the server to see the pages as they look).")
            return out
        return {"view": "none"}
    if ext == ".docx":
        return _docx(path)
    if ext == ".pptx":
        return _pptx(path)
    if ext in (".xlsx", ".xlsm"):
        return _xlsx(path)
    if ext in (".csv", ".tsv"):
        return _csv(path, ext)
    if ext == ".zip":
        return _zip(path)
    raw = path.read_bytes()[: TEXT_MAX * 4 + 4]
    if _binary(raw):
        return {"view": "none"}
    if kind != "doc" and not mime.startswith("text/") and ext not in (".json", ".md", ".markdown", ".ipynb"):
        try:
            raw.decode("utf-8")   # 认不得的扩展名：是 UTF-8 文字才当文字看
        except UnicodeDecodeError:
            return {"view": "none"}
    text = _decode(raw)
    if ext in (".md", ".markdown"):
        return {"view": "doc", "blocks": _chunks(text[:MD_MAX]), "truncated": len(text) > MD_MAX, "label": "Markdown"}
    if ext == ".ipynb":
        try:
            return _notebook(text)
        except (ValueError, AttributeError, TypeError):
            pass
    if ext == ".json" or mime == "application/json":
        try:
            text = json.dumps(json.loads(text), ensure_ascii=False, indent=2)
        except ValueError:
            pass
        return _text(text, True, "JSON")
    if ext in (".html", ".htm") or mime == "text/html":
        return _text(_html_text(text), False, L("网页（只看文字）", "Web page (text only)"))
    mono = ext not in PROSE_EXT
    return _text(text, mono, L("代码", "Code") if mono and ext not in (".log", ".csv") else L("文字", "Text"))


@router.get("/api/files/{fid}/preview")
def preview(fid: str):
    r = _row(fid)
    out = {"ok": True, "file": summary(r)}
    try:
        out.update(build(Path(r["path"]), r["name"], r["mime"] or "", r["kind"], r["text"]))
    except Exception as exc:  # noqa: BLE001 — 预览坏了不影响原件：退回上传时抽的文字，没有就只给文件信息
        note = L(f"预览出错：{str(exc)[:120]}", f"Preview failed: {str(exc)[:120]}")
        out.update(_text(r["text"], False, "") if r["text"] else {"view": "none"}, note=note)
    if out.get("view") == "none" and not out.get("note"):
        out["note"] = L("这种文件在 app 里看不了，可以用浏览器打开。", "This kind of file can't be shown in the app. Try opening it in the browser.")
    return out


@router.get("/api/files/{fid}/page/{n}")
def page_image(fid: str, n: int, w: int = 1200):
    """PDF 这类的第 n 页（从 1 数）渲染成 JPEG。"""
    r = _row(fid)
    path = Path(r["path"])
    w = page_width(w)
    return FileResponse(str(render_page(path, n, w, path.with_name(f"{path.name}.p{n}-{w}.jpg"))), media_type="image/jpeg", headers=CACHE)


def page_width(w: int) -> int:
    return next((x for x in PAGE_WIDTHS if x >= w), PAGE_WIDTHS[-1])


def render_page(path: Path, n: int, w: int, out: Path) -> Path:
    """第 n 页渲染成 JPEG 存到 out（已经有就直接用）。学习台的课件用它：页图存在课程的 .gen 里，不放进课件文件夹。"""
    if out.is_file():
        return out
    m = _mupdf()
    if not m:
        raise HTTPException(404, L("服务器上没装 PyMuPDF", "PyMuPDF isn't installed on the server"))
    with _MU:
        try:
            doc = _open_paged(m, path)
        except Exception as exc:  # noqa: BLE001
            raise HTTPException(404, L("打不开这个文件", "Couldn't open this file")) from exc
        try:
            if doc.needs_pass or not 1 <= n <= doc.page_count:
                raise HTTPException(404, L("没有这一页", "No such page"))
            page = doc[n - 1]
            pw, ph = page.rect.width, page.rect.height
            zoom = w / pw if pw else 1
            if pw * ph * zoom * zoom > PAGE_PIXELS:
                zoom = (PAGE_PIXELS / (pw * ph)) ** 0.5
            data = page.get_pixmap(matrix=m.Matrix(zoom, zoom), alpha=False).tobytes("jpg", jpg_quality=82)
        finally:
            doc.close()
    out.parent.mkdir(parents=True, exist_ok=True)
    tmp = out.with_name(f"{out.name}.{uuid.uuid4().hex[:6]}.tmp")
    tmp.write_bytes(data)
    tmp.replace(out)
    return out
